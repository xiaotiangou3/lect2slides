import os
from typing import List, Dict, Any
from pptx import Presentation
from pptx.util import Inches
from PIL import Image


def create_presentation_from_slides(slides: List[Dict[str, Any]], output_pptx_path: str) -> str:
    """Builds a 16:9 PPTX with full slide images and transcript in Speaker Notes."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    print(f"[Builder] Assembling presentation ({len(slides)} slides)...")
    for item in slides:
        slide = prs.slides.add_slide(blank_layout)
        img_path = item["image_path"]

        if os.path.exists(img_path):
            with Image.open(img_path) as img:
                w, h = img.size
            img_ar = w / h
            target_ar = prs.slide_width / prs.slide_height

            if abs(img_ar - target_ar) < 0.05:
                slide.shapes.add_picture(img_path, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
            elif img_ar > target_ar:
                new_h = int(prs.slide_width / img_ar)
                slide.shapes.add_picture(img_path, Inches(0), int((prs.slide_height - new_h) / 2), prs.slide_width, new_h)
            else:
                new_w = int(prs.slide_height * img_ar)
                slide.shapes.add_picture(img_path, int((prs.slide_width - new_w) / 2), Inches(0), new_w, prs.slide_height)

        slide.notes_slide.notes_text_frame.text = item.get("notes", "")

    os.makedirs(os.path.dirname(os.path.abspath(output_pptx_path)), exist_ok=True)
    prs.save(output_pptx_path)
    return os.path.abspath(output_pptx_path)
